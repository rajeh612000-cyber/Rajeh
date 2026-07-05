#!/usr/bin/env python3
"""Assemble deck/png/*.png into a ready-to-present 16:9 PowerPoint.

Each slide is a single full-bleed branded PNG — text is baked in, so fonts
never substitute and spacing never shifts on another machine.
"""
import glob, os, io, sys
from pptx import Presentation
from pptx.util import Emu
from PIL import Image

# format: 'png' (lossless, exactly as generated — default) or 'jpeg' (lighter, email-safe)
FMT = sys.argv[1] if len(sys.argv) > 1 else 'png'
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PNG = os.path.join(ROOT, 'deck', 'png')
OUT = os.path.join(ROOT, 'Smart_Value_Brand_System.pptx' if FMT == 'png' else 'Smart_Value_Brand_System_light.pptx')

# 16:9 widescreen — 13.333in × 7.5in
W, H = Emu(12192000), Emu(6858000)

prs = Presentation()
prs.slide_width = W
prs.slide_height = H
blank = prs.slide_layouts[6]  # fully blank layout

def as_jpeg(path, q=94):
    """High-quality JPEG (4:4:4, no chroma subsampling) — crisp text, portable size."""
    buf = io.BytesIO()
    Image.open(path).convert('RGB').save(buf, 'JPEG', quality=q, optimize=True, subsampling=0)
    buf.seek(0)
    return buf

files = sorted(glob.glob(os.path.join(PNG, '*.png')))
for f in files:
    slide = prs.slides.add_slide(blank)
    img = f if FMT == 'png' else as_jpeg(f)   # embed the PNG directly, or a JPEG copy
    slide.shapes.add_picture(img, 0, 0, width=W, height=H)

cp = prs.core_properties
cp.title = 'Smart Value™ — Brand System'
cp.author = 'Smart Value™ AI Solutions'
cp.subject = 'Visual Identity Guidelines · Version 1.0 · 2026'
cp.keywords = 'brand, identity, Smart Value, RGM, FMCG'

prs.save(OUT)
size = os.path.getsize(OUT) / (1024 * 1024)
print(f'wrote {os.path.basename(OUT)}  ·  {len(files)} slides  ·  {size:.1f} MB')
